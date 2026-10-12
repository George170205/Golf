import openpyxl
import json
import unicodedata
import re

def normalize(s):
    if not s:
        return ""
    s = str(s).strip().lower()
    s = unicodedata.normalize('NFKD', s).encode('ASCII', 'ignore').decode('utf-8')
    s = re.sub(r'[^a-z0-9]', '', s)
    return s

wb = openpyxl.load_workbook('Copy of LOS SANTOS_2026_27 TOURNAMENT.xlsx', data_only=True)

players_list = []
norm_to_player = {}

ALIASES = {
    normalize("Albert Joel Martínez"): "albert",
    normalize("Albert Joel Martinez"): "albert",
    normalize("Alberto Martinez"): "albert",
    normalize("Albert Martínez"): "albert",
    normalize("Alfonso Noriega"): "alfonsonogueira",
    normalize("Alfonso Nogueira"): "alfonsonogueira",
    normalize("Xavi Vela"): "xaviervela",
    normalize("Xavier Vela"): "xaviervela",
    normalize("Carlos del Río Jr"): "carlosdelriojr",
    normalize("Carlos Del Rio Jr"): "carlosdelriojr",
    normalize("Carlos del Rio Sr"): "carlosdelriosr",
    normalize("Carlos Del Rio"): "carlosdelriosr",
    normalize("Pedro Fernando Gonzalez"): "pedrogonzalez",
    normalize("Pedro Gonzalez"): "pedrogonzalez",
    normalize("Patricio Orozco Loza"): "patricioorozco",
    normalize("Patricio Orozco"): "patricioorozco",
    normalize("Jorge Ojeda Quintana"): "jorgeojeda",
    normalize("Jorge Ojeda"): "jorgeojeda",
    normalize("Mito Camacho"): "miguelcamacho",
    normalize("Miguel Camacho"): "miguelcamacho",
    normalize("José Arturo Sánchez"): "arturosanchez",
    normalize("Jose Arturo Sanchez"): "arturosanchez",
    normalize("Arturo Sanchez"): "arturosanchez",
    normalize("Iván Becerra Serrano"): "ivanbecerra",
    normalize("Ivan Becerra Serrano"): "ivanbecerra",
    normalize("Ivan Becerra"): "ivanbecerra",
}

ws_main = wb['BD_PRINCIPAL']
for r in range(2, ws_main.max_row + 1):
    name = ws_main.cell(r, 2).value
    if not name or str(name).strip() == "":
        continue
    name = str(name).strip()
    ghin = ws_main.cell(r, 3).value
    ghin_str = str(int(ghin)) if isinstance(ghin, (int, float)) and ghin else (str(ghin) if ghin else "")
    hci = ws_main.cell(r, 4).value
    hci_val = float(hci) if isinstance(hci, (int, float)) else 0.0
    cat = ws_main.cell(r, 6).value
    cat_str = str(cat).strip() if cat else "Premiere"
    if "Prem" in cat_str:
        cat_str = "Premiere"
    elif "Elit" in cat_str:
        cat_str = "Elite"
    elif "Adv" in cat_str:
        cat_str = "Advanced"

    p_id = f"p_{len(players_list) + 1}"
    p_obj = {
        "id": p_id,
        "name": name,
        "ghin": ghin_str,
        "hci": round(hci_val, 1),
        "category": cat_str,
        "teamId": None,
        "slot": None
    }
    players_list.append(p_obj)
    
    n_key = normalize(name)
    norm_to_player[n_key] = p_obj
    if n_key in ALIASES:
        norm_to_player[ALIASES[n_key]] = p_obj

def find_player(name_str):
    if not name_str:
        return None
    k = normalize(name_str)
    if k in norm_to_player:
        return norm_to_player[k]
    if k in ALIASES and ALIASES[k] in norm_to_player:
        return norm_to_player[ALIASES[k]]
    for pk, p in norm_to_player.items():
        if k in pk or pk in k:
            return p
    return None

# Assign LIV teams
ws_liv = wb['LIV 2026-27']
teams = [
    {"id": "t_1", "name": "Beerdie Guys", "color": "#1F5A3A"},
    {"id": "t_2", "name": "Slicing Ninjas", "color": "#B8893B"},
    {"id": "t_3", "name": "Chapter 11", "color": "#0F4C75"},
    {"id": "t_4", "name": "Little Balls", "color": "#7E22CE"}
]
team_cols = [
    ("t_1", 1, 2, 3),
    ("t_2", 6, 7, 8),
    ("t_3", 11, 12, 13),
    ("t_4", 16, 17, 18)
]

for team_id, c_team, c_slot, c_player in team_cols:
    for r in range(2, ws_liv.max_row + 1):
        p_name = ws_liv.cell(r, c_player).value
        slot = ws_liv.cell(r, c_slot).value
        if not p_name:
            continue
        p = find_player(p_name)
        if p:
            p["teamId"] = team_id
            p["slot"] = int(float(slot)) if slot else None
        else:
            cat = ws_liv.cell(r, c_player + 1).value
            cat_str = "Premiere" if "Prem" in str(cat) else ("Elite" if "Elit" in str(cat) else "Advanced")
            new_id = f"p_{len(players_list) + 1}"
            new_p = {
                "id": new_id,
                "name": str(p_name).strip(),
                "ghin": "",
                "hci": 15.0,
                "category": cat_str,
                "teamId": team_id,
                "slot": int(float(slot)) if slot else None
            }
            players_list.append(new_p)
            norm_to_player[normalize(p_name)] = new_p

# J1 scores from Premiere, Elite, Advanced
results = []
ctp_j1_players = set()
for cat_name in ['Premiere', 'Elite', 'Advanced']:
    ws = wb[cat_name]
    for r in range(3, ws.max_row + 1):
        p_name = ws.cell(r, 2).value
        if not p_name:
            continue
        p = find_player(p_name)
        if not p:
            continue
        ctp_val = ws.cell(r, 9).value
        if ctp_val and float(ctp_val) > 0:
            ctp_j1_players.add(p["id"])
        
        s1 = ws.cell(r, 10).value
        u1 = ws.cell(r, 11).value
        pts1 = ws.cell(r, 12).value
        if s1 is not None and float(s1) > 0:
            net1 = float(s1)
            hc_app = round(p["hci"] * 0.8)
            gross1 = int(round(net1 + hc_app))
            results.append({
                "id": f"res_r1_{p['id']}",
                "roundId": "r_1",
                "playerId": p["id"],
                "gross": gross1,
                "net": net1,
                "putts": int(float(u1)) if u1 else 0,
                "points": float(pts1) if pts1 else 0,
                "ctp": p["id"] in ctp_j1_players,
                "hciSnapshot": p["hci"]
            })

# J2 scores from J2_RDM_20_SEP sheet
ws2 = wb['J2_RDM_20_SEP']
for r in range(2, ws2.max_row + 1):
    name = ws2.cell(r, 2).value
    gross = ws2.cell(r, 3).value
    hc = ws2.cell(r, 4).value
    net = ws2.cell(r, 5).value
    putts = ws2.cell(r, 6).value
    cat = ws2.cell(r, 7).value
    pts = ws2.cell(r, 8).value
    
    if not name or gross is None:
        continue
    p = find_player(name)
    if not p:
        # Player in J2 who wasn't registered yet?
        p_id = f"p_{len(players_list) + 1}"
        p = {
            "id": p_id,
            "name": str(name).strip(),
            "ghin": "",
            "hci": float(hc) if hc is not None else 15.0,
            "category": str(cat).strip() if cat in ['Premiere', 'Elite', 'Advanced'] else "Elite",
            "teamId": None,
            "slot": None
        }
        players_list.append(p)
        norm_to_player[normalize(name)] = p
        print(f"Added J2 participant to roster: {name}")
    
    # check if already added
    res_id = f"res_r2_{p['id']}"
    if not any(x["id"] == res_id for x in results):
        results.append({
            "id": res_id,
            "roundId": "r_2",
            "playerId": p["id"],
            "gross": int(float(gross)),
            "net": float(net) if net is not None else (float(gross) - round(p["hci"] * 0.8)),
            "putts": int(float(putts)) if putts is not None else 0,
            "points": float(pts) if pts is not None else 0,
            "ctp": False,
            "hciSnapshot": float(hc) if hc is not None else p["hci"]
        })

print(f"Final Count -> Players: {len(players_list)}, Results: {len(results)}")

rounds = [
    {"id": "r_1", "number": 1, "name": "Jornada 1", "course": "Baja Country Club", "shortCourse": "BCC", "date": "2026-08-16", "status": "completed"},
    {"id": "r_2", "number": 2, "name": "Jornada 2", "course": "Real del Mar", "shortCourse": "RDM", "date": "2026-09-20", "status": "completed"},
    {"id": "r_3", "number": 3, "name": "Jornada 3", "course": "Club Campestre Tijuana", "shortCourse": "COR", "date": "2026-10-18", "status": "scheduled"},
    {"id": "r_4", "number": 4, "name": "Jornada 4", "course": "Bajamar Ocean Front", "shortCourse": "BAJ", "date": "2026-11-15", "status": "scheduled"},
    {"id": "r_5", "number": 5, "name": "Jornada 5", "course": "Club Campestre Tijuana", "shortCourse": "COR", "date": "2026-12-13", "status": "scheduled"}
]

points_rules = [
    {"place": 1, "points": 15},
    {"place": 2, "points": 13},
    {"place": 3, "points": 12},
    {"place": 4, "points": 10},
    {"place": 5, "points": 9},
    {"place": 6, "points": 8},
    {"place": 7, "points": 7},
    {"place": 8, "points": 6},
    {"place": 9, "points": 5},
    {"place": 10, "points": 4},
    {"place": 11, "points": 3},
    {"place": 12, "points": 3},
    {"place": 13, "points": 3},
    {"place": 14, "points": 3},
    {"place": 15, "points": 3},
    {"place": 16, "points": 3},
    {"place": 17, "points": 3},
    {"place": 18, "points": 3}
]

ctp_rules = [
    {"place": 1, "points": 10},
    {"place": 2, "points": 7},
    {"place": 3, "points": 5},
    {"place": 4, "points": 3}
]

liv_rules = [
    {"place": 1, "points": 10},
    {"place": 2, "points": 7},
    {"place": 3, "points": 5},
    {"place": 4, "points": 3}
]

tournaments = [
    {
        "id": "tour_1",
        "name": "Torneo Los Santos 2026/27",
        "season": "2026-2027",
        "club": "Los Santos Golf Club",
        "description": "Campeonato Oficial con clasificación individual por categoría (Premiere, Elite, Advanced) y Copa LIV por Equipos.",
        "useHci80": True,
        "tiebreak": "putts",
        "status": "in_progress",
        "createdAt": "2026-08-01"
    }
]

seed_data = {
    "tournaments": tournaments,
    "activeTournamentId": "tour_1",
    "players": players_list,
    "teams": teams,
    "rounds": rounds,
    "results": results,
    "pointsRules": points_rules,
    "ctpRules": ctp_rules,
    "livRules": liv_rules,
    "settings": {
        "useHci80": True,
        "tiebreak": "putts",
        "ctpBonusPoints": 1,
        "theme": "dark"
    }
}

with open("seed_tournament_data.json", "w", encoding="utf-8") as f:
    json.dump(seed_data, f, ensure_ascii=False, indent=2)

print("Exported clean seed_tournament_data.json with ALL real data successfully!")
